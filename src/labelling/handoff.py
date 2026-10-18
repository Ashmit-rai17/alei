"""Labeler hand-off: split labeling_sample.xlsx into data/prashant/.

phase1_read.xlsx   15 rows (5 per claim_type) read cold — no ground_span shown
phase2_label.xlsx  the other 64 rows, full labeling columns
candidate_grounds.csv  phrase-frequency table for freezing the taxonomy
                   (owner's use only; NOT included in the zip)
README.md          written by hand, kept in the folder

Reads only data/labeling_sample.xlsx, which is built from awards_safe.parquet.
"""

import os
import sys
import zipfile

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SRC_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path[:0] = [os.path.join(SRC_DIR, p) for p in ("acquisition", "parsing", "privacy", "labelling", "modelling", "reporting")]
from inventory import DATA  # noqa: E402

SRC = os.path.join(DATA, "labeling_sample.xlsx")
OUT = os.path.join(DATA, "prashant")
ZIP = os.path.join(DATA, "prashant.zip")
PER_TYPE = 5
TYPES = ["life_death", "health", "life_misc"]

P1_COLS = ["award_id", "claim_type", "insurer", "award_text", "reason_in_my_words", "confidence", "notes"]
P2_COLS = ["award_id", "claim_type", "insurer", "award_date", "ground_span", "award_text",
           "ground_1", "ground_2", "ground_3", "ground_4", "ground_5", "corrected", "notes"]

P1_INSTRUCTIONS = [
    "Read each award and write, in your own words, every reason the INSURER gave for rejecting "
    "the claim — not the Ombudsman's view of it, and not who won. More than one reason is normal.",
    "confidence: high / medium / low — how sure you are that you found the insurer's reason(s).",
    "notes: anything odd — the text is cut off, two awards seem joined, a real name slipped "
    "through, you can't find a reason at all.",
    "See README.md for the full brief.",
]
P2_EXTRA = ("The category list for ground_1 … ground_5 is still being frozen. It will arrive "
            "before this phase starts — don't begin phase 2 until you have it.")


def style(ws, widths, fill_from):
    for j, w in enumerate(widths, 1):
        cell = ws.cell(row=1, column=j)
        ws.column_dimensions[cell.column_letter].width = w
        cell.font = Font(bold=True)
        if j >= fill_from:
            cell.fill = PatternFill("solid", fgColor="FFF2CC")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


def write_instructions(xw, lines):
    pd.DataFrame({"instructions": lines}).to_excel(xw, sheet_name="instructions", index=False)
    wi = xw.sheets["instructions"]
    wi.column_dimensions["A"].width = 120
    for row in wi.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def handoff():
    os.makedirs(OUT, exist_ok=True)
    book = pd.read_excel(SRC, sheet_name=None)
    lab = book["labeling"].fillna("")

    p1_idx = [i for t in TYPES for i in lab.index[lab.claim_type == t][:PER_TYPE]]
    p1 = lab.loc[p1_idx].copy()
    p2 = lab.drop(index=p1_idx).copy()
    assert len(p1) == 15 and len(p2) == 64 and not set(p1.award_id) & set(p2.award_id)

    for c in ["reason_in_my_words", "confidence", "notes"]:
        p1[c] = ""
    p1 = p1[P1_COLS]
    p2 = p2[P2_COLS]

    with pd.ExcelWriter(os.path.join(OUT, "phase1_read.xlsx"), engine="openpyxl") as xw:
        p1.to_excel(xw, sheet_name="read", index=False)
        ws = xw.sheets["read"]
        style(ws, [24, 14, 18, 110, 60, 12, 40], fill_from=5)
        dv = DataValidation(type="list", formula1='"high,medium,low"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"F2:F{len(p1) + 1}")
        write_instructions(xw, P1_INSTRUCTIONS)

    instr = [x for x in book["instructions"]["instructions"].fillna("").tolist()
             if not x.startswith(("Sample:", "Drawn per claim_type"))]
    with pd.ExcelWriter(os.path.join(OUT, "phase2_label.xlsx"), engine="openpyxl") as xw:
        p2.to_excel(xw, sheet_name="labeling", index=False)
        ws = xw.sheets["labeling"]
        style(ws, [24, 14, 18, 12, 60, 100, 30, 30, 30, 30, 30, 11, 40], fill_from=7)
        dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"L2:L{len(p2) + 1}")
        write_instructions(xw, [P2_EXTRA, ""] + instr)

    book["candidate_grounds"].to_csv(os.path.join(OUT, "candidate_grounds.csv"), index=False)

    # the zip is what gets sent: phase files + README only
    send = ["README.md", "phase1_read.xlsx", "phase2_label.xlsx"]
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for f in send:
            z.write(os.path.join(OUT, f), arcname=f"prashant/{f}")

    print(f"phase1: {len(p1)} rows {p1.claim_type.value_counts().to_dict()}")
    print(f"phase2: {len(p2)} rows {p2.claim_type.value_counts().to_dict()}")
    print(f"zip -> {ZIP}: {send}  (candidate_grounds.csv kept out)")


if __name__ == "__main__":
    handoff()
